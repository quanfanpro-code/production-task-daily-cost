"""行为验收：使用合成资料，不包含真实生产数据。"""
import datetime as dt
import importlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))


def record(process='精磨', date='2026-01-03', kind='point', **kw):
    return dict(process=process, date=date, kind=kind,
                source={'file': '合成单据.pdf', 'page': 2, 'location': '第1行'}, **kw)


def task(records, start='2026-01-01', **kw):
    return dict(task_no='SAMPLE-001', product_name='示例零件', spec='10×20', material='石英',
                start={'date': start, 'basis': '生产计划领料'}, records=records, **kw)


class Rules(unittest.TestCase):
    def run_task(self, value):
        self.assertTrue((ROOT / 'scripts/日期推算.py').exists(), '日期推算尚未实现')
        return importlib.import_module('日期推算').build([value])

    def test_跨月逐日且不设14天截断(self):
        out = self.run_task(task([record(date='2026-02-02')], start='2026-01-10'))
        self.assertEqual(len(out['normal']), 24)
        self.assertEqual(out['normal'][0]['date'], '2026-01-10')
        self.assertEqual(out['normal'][-1]['date'], '2026-02-02')

    def test_工序时点按日期排序且跨工序共享交接日(self):
        out = self.run_task(task([record('双轴', '2026-01-05'), record('精磨', '2026-01-03')]))
        self.assertEqual([(r['process'], r['date']) for r in out['normal']],
                         [('精磨','2026-01-01'),('精磨','2026-01-02'),('精磨','2026-01-03'),
                          ('双轴','2026-01-03'),('双轴','2026-01-04'),('双轴','2026-01-05')])

    def test_同工序多次记录合并同日(self):
        out = self.run_task(task([record(date='2026-01-03'),record(date='2026-01-05'),record(date='2026-01-05')]))
        self.assertEqual(len(out['normal']), 5)

    def test_明确同日起止不能回推(self):
        out = self.run_task(task([record(date='2026-01-20',kind='range',end='2026-01-20')]))
        self.assertEqual([r['date'] for r in out['normal']], ['2026-01-20'])

    def test_实际明确区间允许不同工序重叠(self):
        out = self.run_task(task([record(kind='range',end='2026-01-05'),
                                  record('双轴','2026-01-04','range',end='2026-01-06')]))
        self.assertEqual(len(out['normal']), 6)

    def test_同日不同工序分别计入(self):
        out = self.run_task(task([record('双轴'),record('精磨')]))
        self.assertEqual(len(out['normal']), 4)
        self.assertEqual({r['process'] for r in out['normal'] if r['date']=='2026-01-03'}, {'精磨','双轴'})

    def test_计划补齐且实际优先(self):
        out = self.run_task(task([record()], plan=[{'process':'精磨','start':'2026-01-01','end':'2026-01-10'},
                                               {'process':'双轴','start':'2026-01-04','end':'2026-01-06'}]))
        self.assertEqual(len(out['normal']), 6)
        self.assertNotIn('2026-01-10', [r['date'] for r in out['normal']])

    def test_检验无日期采用交期并顺延(self):
        out = self.run_task(task([record(), record('成品检验',None,'inspection')],delivery_date='2026-01-02'))
        self.assertEqual([(r['process'],r['date']) for r in out['normal'] if r['process']=='成品检验'], [('成品检验','2026-01-03')])

    def test_有记录缺日期必须补齐(self):
        out = self.run_task(task([record(date=None)],delivery_date='2026-01-05'))
        self.assertTrue(out['normal'])
        self.assertTrue(any('推算' in r.get('nature','') for r in out['trace']))

    def test_空白任务仍留痕(self):
        out = self.run_task(task([]))
        self.assertEqual(out['normal'], [])
        self.assertTrue(out['trace'])
        self.assertEqual(out['trace'][0]['task_no'], 'SAMPLE-001')

    def test_异常183天边界整任务互斥分流(self):
        for end,count,abnormal in [('2026-07-01',182,False),('2026-07-02',183,True),('2026-07-03',184,True)]:
            with self.subTest(end=end):
                out=self.run_task(task([record(date=end),record('成品检验',end,'inspection')]))
                active=out['abnormal'] if abnormal else out['normal']
                self.assertEqual(len(active),count+1)
                self.assertEqual(out['normal'] if abnormal else out['abnormal'],[])

    def test_并行行数不能当工期(self):
        out=self.run_task(task([record(kind='range',date='2026-01-01',end='2026-04-10'),
                                record('双轴','2026-01-01','range',end='2026-04-10')]))
        self.assertEqual(len(out['normal']),200)
        self.assertEqual(out['abnormal'],[])

    def test_输入反序结果一致(self):
        rs=[record('精磨','2026-01-03'),record('双轴','2026-01-03'),record('精磨','2026-01-05')]
        self.assertEqual(self.run_task(task(rs)),self.run_task(task(list(reversed(rs)))))

    def test_实写和时点插值留痕不同(self):
        out=self.run_task(task([record()]))
        traces={r['date']:r for r in out['trace'] if r.get('date')}
        self.assertEqual(traces['2026-01-01']['nature'],'实际时点连续推算')
        self.assertEqual(traces['2026-01-03']['nature'],'实际完成时点')

    def test_跨年区间(self):
        out=self.run_task(task([record(date='2026-01-02')],start='2025-12-30'))
        self.assertEqual([r['date'] for r in out['normal']],['2025-12-30','2025-12-31','2026-01-01','2026-01-02'])

    def test_仅计划检验取完成当天(self):
        out=self.run_task(task([],plan=[{'process':'成品检验','start':'2026-01-05','end':'2026-01-08'}]))
        self.assertEqual([(r['process'],r['date']) for r in out['normal']],[('成品检验','2026-01-08')])

    def test_有计划工序缺日期仍补齐(self):
        out=self.run_task(task([],plan=[{'process':'精磨','raw':'加工日期模糊'}]))
        self.assertEqual(len(out['normal']),1)

    def test_候选循环胜负不按输入顺序选择(self):
        engine=importlib.import_module('日期推算')
        cs=[{'value':'2026-01-03','beats':['2026-01-04']},{'value':'2026-01-04','beats':['2026-01-05']},{'value':'2026-01-05','beats':['2026-01-03']}]
        self.assertIsNone(engine.choose_candidate(cs))
        self.assertIsNone(engine.choose_candidate(list(reversed(cs))))

    def test_补齐后起点不能晚于完成(self):
        out=self.run_task(task([record(date=None)],delivery_date='2025-12-30'))
        self.assertGreaterEqual(out['tasks'][0]['span_days'],1)

    def test_只有实际检验日期也能交付(self):
        value=task([record('成品检验','2026-02-02','inspection')]);value.pop('start')
        out=self.run_task(value)
        self.assertEqual([r['date'] for r in out['normal']],['2026-02-02'])

    def test_只有计划完成日也能交付(self):
        value=task([],plan=[{'process':'精磨','end':'2026-02-02'}]);value.pop('start')
        self.assertEqual(len(self.run_task(value)['normal']),1)


class Files(unittest.TestCase):
    def test_文件备份与路径防越界(self):
        from 文件保护 import backup,write_bytes,safe_child,digest
        base=Path(tempfile.mkdtemp(prefix='文件保护测试_'));path=base/'原文件.txt';path.write_bytes(b'original')
        copy=backup(path);self.assertEqual(digest(path),digest(copy))
        write_bytes(path,b'updated');self.assertEqual(copy.read_bytes(),b'original')
        with self.assertRaises(ValueError):safe_child(base,'../outside.pdf')

    def test_重复件可追溯且来源变动拒绝(self):
        from 预处理拆分 import split
        from 文件保护 import digest
        import pymupdf
        base=Path(tempfile.mkdtemp(prefix='重复件测试_'));src=base/'源';src.mkdir()
        doc=pymupdf.open();doc.new_page();doc.new_page();doc.save(src/'两页.pdf');doc.close()
        plan={'source_root':str(src),'files':{'两页.pdf':{'pages':2,'sha256':digest(src/'两页.pdf')}},'segments':[
            {'source_file':'两页.pdf','start':1,'end':1,'task_no':'A','id':'keep','action':'keep'},
            {'source_file':'两页.pdf','start':2,'end':2,'task_no':'A','action':'duplicate','duplicate_of':'keep','reason':'编号和填写内容一致，采用清楚来源'}]}
        out=split(plan,base/'输出');self.assertEqual(out['kept_pages'],1);self.assertEqual(len(out['pages']),2)
        self.assertEqual(out['pages'][1]['output_file'],'A.pdf');self.assertIsNone(out['pages'][1]['output_page'])
        plan['files']['两页.pdf']['sha256']='bad'
        with self.assertRaisesRegex(ValueError,'变化'):split(plan,base/'拒绝')

    def test_拆分页数来源和漏页拒绝(self):
        self.assertTrue((ROOT/'scripts/预处理拆分.py').exists(),'拆分功能尚未实现')
        mod=importlib.import_module('预处理拆分')
        import pymupdf
        base=Path(tempfile.mkdtemp(prefix='生产任务测试_'))
        src=base/'原件';src.mkdir()
        doc=pymupdf.open()
        for s in ['TASK-A','DRAWING','TASK-B']:
            p=doc.new_page();p.insert_text((70,70),s)
        doc.save(src/'合成.pdf');doc.close()
        import hashlib
        digest=hashlib.sha256((src/'合成.pdf').read_bytes()).hexdigest()
        segments=[dict(source_file='合成.pdf',start=1,end=2,task_no='A',action='keep'),
                  dict(source_file='合成.pdf',start=3,end=3,task_no='B',action='keep')]
        plan={'source_root':str(src),'files':{'合成.pdf':{'pages':3,'sha256':digest}},'segments':segments}
        result=mod.split(plan,base/'拆分')
        self.assertEqual(len(result['pages']),3)
        self.assertEqual(hashlib.sha256((src/'合成.pdf').read_bytes()).hexdigest(),digest)
        self.assertEqual(sorted(pymupdf.open(p).page_count for p in (base/'拆分').glob('*.pdf')),[1,2])
        plan['segments']=segments[:1]
        with self.assertRaisesRegex(ValueError,'页'):
            mod.split(plan,base/'错误输出')

    def test_生成真日期六列和完整留痕(self):
        self.assertTrue((ROOT/'scripts/生成交付.py').exists(),'交付生成尚未实现')
        engine=importlib.import_module('日期推算'); export=importlib.import_module('生成交付')
        out=Path(tempfile.mkdtemp(prefix='生产任务表格测试_'))
        value=task([record()],product_override='=1+1')
        value['product_name']=value.pop('product_override')
        result=engine.build([value])
        path=export.write_delivery(result,out)
        from openpyxl import load_workbook
        wb=load_workbook(path)
        self.assertEqual(wb.sheetnames,['逐日明细','异常工期逐日明细','留痕'])
        ws=wb['逐日明细']
        self.assertEqual(ws.max_column,6)
        self.assertEqual([c.value for c in ws[1]],['任务编号','产品名称','规格','材料','工序','日期'])
        self.assertIsInstance(ws['F2'].value,dt.datetime)
        self.assertEqual(ws['B2'].data_type,'s')
        self.assertEqual(ws['B2'].value,'=1+1')
        self.assertEqual(export.verify_workbook(path)['normal_rows'],3)
        wb.close()


if __name__=='__main__':
    unittest.main(verbosity=2)
