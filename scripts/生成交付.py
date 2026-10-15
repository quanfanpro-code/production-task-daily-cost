"""生成三页工作簿和中文说明；回读验证日记录与来源。"""
import argparse
from datetime import datetime
import io
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from 文件保护 import write_bytes,write_json,write_text
from 日期推算 import build

HEADERS=['任务编号','产品名称','规格','材料','工序','日期']
KEYS=['task_no','product_name','spec','material','process','date']
TRACE=['任务编号','工序','日期','记录性质','判断理由','采用起点','起点依据','任务日历跨度','异常长工期','原文与全部证据']


def verify_workbook(path):
    from openpyxl import load_workbook
    # 明确检查可忽略命名空间前缀已声明，避免Excel打开时报修复。
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.endswith('.xml'):
                raw=z.read(name)
                namespaces={prefix for _,(prefix,_) in ET.iterparse(io.BytesIO(raw),events=['start-ns'])}
                for element in ET.fromstring(raw).iter():
                    for prefix in element.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable','').split():
                        if prefix not in namespaces:raise ValueError('未声明命名空间：'+name+' '+prefix)
    wb=load_workbook(path,read_only=True,data_only=False)
    try:
        if wb.sheetnames!=['逐日明细','异常工期逐日明细','留痕']:
            raise ValueError('工作表名称不符合契约')
        seen=set();task_sets=[];counts=[]
        for name in wb.sheetnames[:2]:
            ws=wb[name]
            it=ws.iter_rows(values_only=True)
            if list(next(it))!=HEADERS:raise ValueError(name+'必须严格6列')
            tasks=set();count=0
            for row in it:
                if len(row)!=6 or not isinstance(row[5],datetime):raise ValueError('日期不是Excel日期')
                key=(row[0],row[4],row[5])
                if key in seen:raise ValueError('发现重复逐日行')
                seen.add(key);tasks.add(row[0]);count+=1
            task_sets.append(tasks);counts.append(count)
        if task_sets[0]&task_sets[1]:raise ValueError('同任务同时进入正常和异常页')
        tr=list(wb['留痕'].iter_rows(values_only=True))[1:]
        traced={r[0] for r in tr}
        if not set.union(*task_sets)<=traced:raise ValueError('日记录没有对应任务留痕')
        return {'normal_rows':counts[0],'abnormal_rows':counts[1],'trace_rows':len(tr),
                'normal_tasks':len(task_sets[0]),'abnormal_tasks':len(task_sets[1]),'traced_tasks':len(traced),
                'six_columns':True,'true_dates':True,'unique_rows':True,'namespace_check':True}
    finally:wb.close()


def write_delivery(result,output):
    import xlsxwriter
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    buf=io.BytesIO()
    wb=xlsxwriter.Workbook(buf,{'in_memory':True,'strings_to_formulas':False,'strings_to_urls':False})
    header=wb.add_format({'bold':True,'bg_color':'#DDEBF7','border':1})
    date_fmt=wb.add_format({'num_format':'yyyy-mm-dd'})
    cell=wb.add_format({'font_size':10,'valign':'top'})
    date_fmt.set_font_size(10)
    for name,rows in [('逐日明细',result['normal']),('异常工期逐日明细',result['abnormal'])]:
        ws=wb.add_worksheet(name);ws.freeze_panes(1,0);ws.set_default_row(16)
        ws.write_row(0,0,HEADERS,header)
        for i,row in enumerate(rows,1):
            for j,key in enumerate(KEYS):
                if key=='date':ws.write_datetime(i,j,datetime.fromisoformat(row[key]),date_fmt)
                else:ws.write_string(i,j,str(row.get(key) or ''),cell)
        ws.set_column(0,0,27);ws.set_column(1,1,24);ws.set_column(2,2,24)
        ws.set_column(3,4,12);ws.set_column(5,5,12)
        ws.autofilter(0,0,len(rows),5)
    ws=wb.add_worksheet('留痕');ws.write_row(0,0,TRACE,header);ws.freeze_panes(1,0)
    ws.set_column(0,0,27);ws.set_column(1,3,16);ws.set_column(4,4,60)
    ws.set_column(5,6,22);ws.set_column(7,8,16);ws.set_column(9,9,90)
    index=1
    for r in result['trace']:
        evidence=json.dumps(r.get('evidence',{}),ensure_ascii=False,sort_keys=True)
        # 长证据分行保存，禁止Excel单元格上限造成静默截断。
        chunks=[evidence[i:i+30000] for i in range(0,len(evidence),30000)] or ['']
        for n,chunk in enumerate(chunks):
            values=[r['task_no'],r.get('process',''),r.get('date',''),r.get('nature',''),
                    r.get('reason','')+(f'；证据分段{n+1}/{len(chunks)}' if len(chunks)>1 else ''),
                    r.get('start',''),r.get('start_basis',''),r.get('span_days',0),'是' if r.get('abnormal') else '否',chunk]
            for j,value in enumerate(values):
                if j in (2,5) and value:ws.write_datetime(index,j,datetime.fromisoformat(value),date_fmt)
                elif j==7:ws.write_number(index,j,value,cell)
                else:ws.write_string(index,j,str(value),cell)
            index+=1
    ws.autofilter(0,0,max(1,index-1),9)
    wb.close()
    path=write_bytes(output/'生产任务_成本逐日明细.xlsx',buf.getvalue())
    check=verify_workbook(path)
    if check['normal_rows']!=len(result['normal']) or check['abnormal_rows']!=len(result['abnormal']):
        raise ValueError('写入回读行数不一致')
    if check['traced_tasks']!=len(result['tasks']):raise ValueError('并非所有任务已写入留痕')
    write_json(output/'完整计算结果.json',result)
    abnormal=[t for t in result['tasks'] if t['abnormal']]
    blank=[t for t in result['tasks'] if not t['row_count']]
    lines=['# 生产任务特殊情况汇报','',f"规则版本：{result['rule_version']}。",'',
           f"处理任务{len(result['tasks'])}项；正常逐日{len(result['normal'])}行；异常逐日{len(result['abnormal'])}行。",
           '正常与异常工作表互斥；异常任务保留完整日期，没有截断。','',
           '## 异常长工期','', '任务起点至最后加工完成，首尾均计，183天及以上归异常页。']
    lines += [f"- {t['task_no']}：{t['start']}至{t['last_process_date']}，{t['span_days']}天；起点依据：{t['start_basis']}。" for t in abnormal] or ['本批无。']
    lines += ['','## 无逐日记录的任务','']
    lines += [f"- {t['task_no']}：无加工/检验发生依据，原始资料和空白事实已在留痕登记。" for t in blank] or ['本批无。']
    lines += ['','## 其他已处理情况','']
    lines += [f"- {w['task_no']}：{w['message']}" for w in result['warnings']] or ['无额外事项。']
    write_text(output/'特殊情况汇报.md','\n'.join(lines)+'\n')
    return path


if __name__=='__main__':
    ap=argparse.ArgumentParser(description='从核定任务数据生成6列逐日明细、异常页与留痕')
    ap.add_argument('--input',required=True);ap.add_argument('--output',required=True)
    ap.add_argument('--pages',help=argparse.SUPPRESS)
    args=ap.parse_args()
    data=json.loads(Path(args.input).read_text(encoding='utf-8-sig'))
    result=build(data['tasks'] if isinstance(data,dict) else data)
    path=write_delivery(result,args.output)
    print(json.dumps({'workbook':str(path.resolve()),'verified':True,'tasks':len(result['tasks']),
                      'normal_rows':len(result['normal']),'abnormal_rows':len(result['abnormal']),
                      'abnormal_tasks':[t for t in result['tasks'] if t['abnormal']],
                      'blank_tasks':[t['task_no'] for t in result['tasks'] if not t['row_count']],
                      'warnings':result['warnings'],'date_checks':result['date_checks']},ensure_ascii=False))
